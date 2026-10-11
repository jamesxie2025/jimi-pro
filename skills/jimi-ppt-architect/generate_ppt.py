import sys
import json
from pptx import Presentation
from pptx.util import Pt

def generate_ppt(json_file, output_file, template_file=None):
    # 1. 读取结构化的 JSON 大纲
    with open(json_file, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 2. 加载用户提供的高级模板，或者使用默认空白模板
    if template_file:
        prs = Presentation(template_file)
        print(f"Using template: {template_file}")
    else:
        prs = Presentation()
        print("Using default blank presentation.")

    # 3. 创建封面页 (Title Slide)
    # 通常模板的第一种布局（索引0）为标题页
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = data.get('title', 'Untitled Presentation')
    subtitle.text = data.get('subtitle', '')

    # 4. 创建正文内容页 (Content Slides)
    # 通常模板的第二种布局（索引1）为标题与内容页
    bullet_slide_layout = prs.slide_layouts[1]
    
    for slide_data in data.get('slides', []):
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        # 设置幻灯片标题
        title_shape.text = slide_data.get('title', '')
        
        # 填充要点列表
        tf = body_shape.text_frame
        tf.clear() # 清空模板中可能存在的默认占位文字
        
        for point in slide_data.get('content', []):
            p = tf.add_paragraph()
            p.text = str(point)
            p.level = 0
            
    # 5. 保存生成的 PPT
    prs.save(output_file)
    print(f"✅ Successfully saved presentation to {output_file}")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_ppt.py <input.json> <output.pptx> [template.pptx]")
        sys.exit(1)
        
    json_path = sys.argv[1]
    out_path = sys.argv[2]
    template_path = sys.argv[3] if len(sys.argv) > 3 else None
    
    try:
        generate_ppt(json_path, out_path, template_path)
    except Exception as e:
        print(f"❌ Error generating presentation: {e}")
        sys.exit(1)
